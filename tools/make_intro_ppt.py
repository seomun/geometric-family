# 「기하학 가족」 제품 소개서(PPT) — python tools/make_intro_ppt.py
# 출력: 다운로드 폴더 기하학가족_제품소개서.pptx (문의 메일은 ~/.gf/store.env 에서만 읽어 넣고 저장소에는 넣지 않는다)
import os, re
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import segno

R = os.path.dirname(os.path.dirname(os.path.abspath(__file__))) + '/'
TMP = os.path.join(os.environ.get('TEMP', '/tmp'), 'gf_intro'); os.makedirs(TMP, exist_ok=True)
OUT = os.path.join(os.path.expanduser('~'), 'Downloads', '기하학가족_제품소개서.pptx')
env = os.path.expanduser('~/.gf/store.env')
EMAIL = ''
if os.path.exists(env):
    m = re.search(r'SUPPORT_EMAIL\s*=\s*(\S+)', open(env, encoding='utf-8').read()); EMAIL = m.group(1).strip('"\'') if m else ''
PLAY = 'https://seomun.github.io/geometric-family/play/'

IN = 914400
CREAM = RGBColor(0xFF, 0xF6, 0xEA); INK = RGBColor(0x4A, 0x3A, 0x30); SUB = RGBColor(0x8A, 0x75, 0x66)
NEMO = RGBColor(0xE0, 0x7B, 0x39); SEMO = RGBColor(0x3A, 0x8F, 0xD6); DONG = RGBColor(0x3B, 0x4A, 0x7A)
YEL = RGBColor(0xF6, 0xDE, 0x5C); WHITE = RGBColor(0xFF, 0xFF, 0xFF); LINE = RGBColor(0xE8, 0xD9, 0xC6)
PINK = RGBColor(0xFF, 0xE3, 0xEA); SKY = RGBColor(0xE3, 0xF0, 0xFF)
FONT = '맑은 고딕'
p = Presentation(); p.slide_width = 12192000; p.slide_height = 6858000; BL = p.slide_layouts[6]

def bg(s, c=CREAM):
    f = s.background.fill; f.solid(); f.fore_color.rgb = c

def text(s, x, y, w, h, t, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, lsp=1.15):
    tb = s.shapes.add_textbox(int(x*IN), int(y*IN), int(w*IN), int(h*IN)); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, ln in enumerate(t if isinstance(t, list) else [t]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.alignment = align; para.line_spacing = lsp
        for j, seg in enumerate(ln.split('**')):
            if not seg: continue
            r = para.add_run(); r.text = seg; r.font.size = Pt(size); r.font.name = FONT; r.font.bold = bold or j % 2 == 1; r.font.color.rgb = color

def box(s, x, y, w, h, fill=WHITE, line=LINE, adj=0.1):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(x*IN), int(y*IN), int(w*IN), int(h*IN))
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.shadow.inherit = False; b.adjustments[0] = adj
    if line is None: b.line.fill.background()
    else: b.line.color.rgb = line; b.line.width = Pt(1)
    return b

def crop(path, box_px, name):
    out = os.path.join(TMP, name); Image.open(R + path).crop(box_px).save(out); return out

def img(s, path, x, y, w=None, h=None, fit=None):
    full = path if os.path.isabs(path) else R + path
    if path.startswith('assets/') and path.endswith('.png'):  # 캐릭터 PNG 는 투명 여백을 잘라 크기를 맞춘다
        im = Image.open(full)
        if im.mode == 'RGBA' and im.getbbox():
            full = os.path.join(TMP, path.replace('/', '_')); im.crop(im.split()[-1].getbbox()).save(full)
    iw, ih = Image.open(full).size
    if fit:  # fit=(w,h) 안에 비율 유지, 가운데
        fw, fh = fit; sc = min(fw/iw, fh/ih); w, h = iw*sc, ih*sc; x += (fw-w)/2; y += (fh-h)/2
    elif w and not h: h = w*ih/iw
    elif h and not w: w = h*iw/ih
    return s.shapes.add_picture(full, int(x*IN), int(y*IN), int(w*IN), int(h*IN))

def title(s, kicker, t, sub=None, color=NEMO):
    text(s, 0.6, 0.35, 6, 0.4, kicker, 13, color, True)
    text(s, 0.6, 0.62, 12.2, 0.7, t, 30, INK, True)
    if sub: text(s, 0.6, 1.3, 12.2, 0.4, sub, 15, SUB)

def foot(s): text(s, 0.6, 7.08, 6, 0.3, '기하학 가족 · 제품 소개서', 10, SUB)

FAM3 = ['assets/nemo_dad/good.png', 'assets/wife/good.png', 'assets/dong_dad/good.png']
APPS = [('icon_toddler', '기하학 가족 놀이터', '유아 교육 놀이', '3~6세', 'kid'), ('icon_color', '막둥이 색칠북', '색칠·스티커', '3~7세', 'kid'),
        ('icon_day', '막둥이의 하루', '생활 습관 놀이', '3~6세', 'kid'), ('icon_tables', '세 가족 식탁', '방치형·사연', '성인', 'ad'),
        ('icon_merge', '도형 합치기', '머지 퍼즐', '13세↑', 'ad'), ('icon_quiz', '당신은 어느 도형?', '성격 테스트', '13세↑', 'ad'),
        ('icon_spot', '다른 그림 찾기', '틀린 그림', '13세↑', 'ad'), ('icon_block', '도형 블록', '블록 퍼즐', '13세↑', 'ad'),
        ('icon_sort', '정리의 달인', '분류 퍼즐', '13세↑', 'ad'), ('icon_tile', '세 가족 짝 맞추기', '타일 매치', '13세↑', 'ad')]

# 1 표지
s = p.slides.add_slide(BL); bg(s)
for i, f in enumerate(FAM3): img(s, f, 6.6 + i*2.1, 1.0, fit=(1.95, 2.9))
text(s, 0.8, 1.2, 6, 1.0, '기하학 가족', 60, INK, True)
text(s, 0.8, 2.45, 5.8, 1.2, ['네모·세모·동그라미 세 도형 가족이', '같은 일을 각자 다르게 겪는 이야기'], 19, SUB)
box(s, 0.8, 3.75, 5.2, 0.62, YEL, None); text(s, 0.8, 3.75, 5.2, 0.62, '웹툰 · 블로그 · 쇼츠 · 게임 10종', 18, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
for i, a in enumerate(APPS): img(s, f'games/demo/{a[0]}.png', 0.8 + i*1.17, 5.2, w=1.0)
text(s, 0.8, 6.35, 11, 0.4, '제품 소개서 · 2026년 10월', 13, SUB)

# 2 한 장으로 보는 IP
s = p.slides.add_slide(BL); bg(s); title(s, 'ONE IP', '하나의 가족, 네 개의 창', '어디서 만나도 같은 가족, 같은 질문 — 「당신은 어느 도형인가요?」'); foot(s)
box(s, 4.65, 2.75, 4.0, 2.6, WHITE, None, 0.5)
for i, f in enumerate(FAM3): img(s, f, 4.9 + i*1.22, 3.05, h=1.5)
text(s, 4.65, 4.65, 4.0, 0.5, '기하학 가족', 20, INK, True, PP_ALIGN.CENTER)
quad = [('웹툰', '잔잔한 공감툰', 'content/blog/ep02_cuts/c01.png', 0.6, 1.95, NEMO), ('블로그', '네이버 웹툰 연재 · 티스토리 정보', 'content/blog/intro_cuts/c01.png', 9.0, 1.95, SEMO),
        ('쇼츠', '통계 + 세 가족 반응', 'content/shorts/s03/f3.png', 0.6, 4.55, DONG), ('게임 10종', '하나의 「우리 집」 공유', 'notes/snapshots/saga_block_start.png', 9.0, 4.55, NEMO)]
for nm, d, im, x, y, c in quad:
    box(s, x, y, 3.7, 2.35)
    img(s, im, x+0.15, y+0.15, fit=(1.5, 2.05))
    text(s, x+1.8, y+0.55, 1.8, 0.5, nm, 20, c, True)
    text(s, x+1.8, y+1.1, 1.8, 1.0, d, 13, INK)

# 3 캐릭터
s = p.slides.add_slide(BL); bg(s); title(s, 'CHARACTERS', '세 가족을 소개합니다', '고등학교 동창 세 남자, 그리고 그 가족들'); foot(s)
fams = [('네모네', '아이 넷, 할머니까지 일곱 식구. 많고 벅차지만 북적이는 행복', NEMO,
         ['assets/nemo_dad/good.png', 'assets/nemo_mom/good.png', 'assets/nemo_grandma/good.png', 'assets/nemo_kids/kid1.png', 'assets/nemo_kids/kid2.png', 'assets/nemo_kids/kid3.png', 'assets/nemo_kids/baby.png']),
        ('세모네', '아이 없는 맞벌이 부부. 티격태격과 달달함 사이의 극단', SEMO, ['assets/husband/good.png', 'assets/wife/good.png']),
        ('동그라미네', '완벽한데 어딘가 비어 있는 집. 정확하고 조용한 사람들', DONG, ['assets/dong_dad/good.png'])]
y = 1.95
for nm, d, c, ims in fams:
    box(s, 0.6, y, 12.1, 1.55)
    text(s, 0.85, y+0.25, 1.9, 0.5, nm, 22, c, True); text(s, 0.85, y+0.8, 3.0, 0.7, d, 12.5, INK)
    for k, f in enumerate(ims): img(s, f, 4.1 + k*1.2, y+0.1, fit=(1.15, 1.35))
    y += 1.68

# 4 세계관
s = p.slides.add_slide(BL); bg(s); title(s, 'WORLD', '하나의 사건, 세 가족의 반응', '같은 부고를 받아도, 같은 저녁을 먹어도 — 세 집은 다르게 겪는다'); foot(s)
rows = [('같은 부고를 받았다면', ['종일 고민한다', '대놓고 물어본다', '미리 준비한다']), ('같은 저녁 시간', ['왁자지껄 혼자가 소원', '여행 가서 둘이', '고급 식탁, 빈 의자']), ('퇴근길 최악의 날', ['막내가 와락 안기면 끝', '술 한잔, 세계 일주 상상', '아무도 기다리지 않는 집'])]
for k, f in enumerate(FAM3): img(s, f, 4.55 + k*2.75, 1.85, h=1.15)
for i, (q, ans) in enumerate(rows):
    y = 3.15 + i*1.15
    box(s, 0.6, y, 3.7, 0.95, YEL, None); text(s, 0.6, y, 3.7, 0.95, q, 15, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    for k, a in enumerate(ans):
        box(s, 4.45 + k*2.75, y, 2.6, 0.95); text(s, 4.45 + k*2.75, y, 2.6, 0.95, a, 13.5, [NEMO, SEMO, DONG][k], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.6, 12.1, 0.4, '행복이 먼저, 아이러니는 뒤에 · 결말은 설명하지 않는다 · 매 화 끝 「당신은 어느 도형인가요?」', 13, SUB, True, PP_ALIGN.CENTER)

# 5 웹툰
s = p.slides.add_slide(BL); bg(s); title(s, 'WEBTOON', '웹툰 — 50대의 잔잔한 공감툰', '1화 「단톡방」 · 2화 「부고」 · 원고 15편 · 한 화 10컷', NEMO); foot(s)
for k, c in enumerate(['c01', 'c03', 'c05', 'c08']): img(s, f'content/blog/ep02_cuts/{c}.png', 0.6 + k*3.05, 1.95, fit=(2.9, 2.5))
for k, c in enumerate(['c01', 'c06', 'c08', 'c09']): img(s, f'content/blog/ep01_short_cuts/{c}.png', 0.6 + k*3.05, 4.55, fit=(2.9, 2.4))

# 6 블로그
s = p.slides.add_slide(BL); bg(s); title(s, 'BLOG', '블로그 — 두 개의 목소리', '같은 세계, 다른 글 (검색 중복 없이)', SEMO); foot(s)
for k, c in enumerate(['c01', 'c02', 'c03']): img(s, f'content/blog/intro_cuts/{c}.png', 0.6 + k*2.05, 1.95, fit=(1.95, 2.45))
box(s, 0.6, 4.55, 6.0, 2.3)
text(s, 0.85, 4.75, 5.6, 0.5, '네이버 블로그 — 웹툰 연재', 19, SEMO, True)
text(s, 0.85, 5.3, 5.6, 1.5, ['· 화자: 세모(별명) — 댓글로 소통', '· 웹툰 본편 + 컷 사이 짧은 글', '· blog.naver.com/gihahakf'], 13.5, INK, lsp=1.3)
box(s, 6.9, 1.95, 5.8, 4.9)
text(s, 7.15, 2.15, 5.3, 0.5, '티스토리 — 「짠내나는 경제학」', 19, NEMO, True)
img(s, 'assets/nemo_dad/worry.png', 10.6, 2.75, h=1.8)
text(s, 7.15, 2.8, 3.4, 3.6, ['· 화자: 네모(작가 자신)', '· 퇴직 후 혼자 버티는 50대의 생활 경제', '· 통계·생활 정보 — 검색 유입', '· 웹툰 소재와 이어지는 정보 글(부의금·명절·건강검진·퇴직)'], 13.5, INK, lsp=1.35)

# 7 쇼츠
s = p.slides.add_slide(BL); bg(s); title(s, 'SHORTS', '쇼츠 — 통계 한 줄, 세 가족 반응', '작가 직접 나레이션 · 큰 글씨 · 움직이는 컷 · 유튜브·릴스·틱톡', DONG); foot(s)
for k, f in enumerate(['content/shorts/s03/f1.png', 'content/shorts/s03/f2.png', 'content/shorts/s03/f3.png', 'content/shorts/s03/f4.png', 'content/shorts/s03/f5.png', 'content/shorts/s02/f1.png']):
    img(s, f, 0.6 + k*2.05, 2.0, fit=(1.9, 3.4))
text(s, 0.6, 5.65, 12.1, 1.1, ['「아이랑 저녁, 매일 같이 드세요?」 — 초6 땐 셋 중 둘, 고3이 되면 다섯 중 하나 (질병관리청 조사)', '「부의금, 얼마 하세요?」 · 「코스피 3천, 우리 집은?」 — 이슈 + 통계 + 효능감'], 14, INK, lsp=1.35)

# 8 게임 라인업
s = p.slides.add_slide(BL); bg(s); title(s, 'GAMES', '게임 10종 — 장르는 열, 집은 하나', '유아 3종 + 성인 7종 · 출시 후보 완성 · 웹에서 바로 체험'); foot(s)
for i, (ic, nm, g, age, grp) in enumerate(APPS):
    x = 0.6 + (i % 5)*2.45; y = 1.95 + (i//5)*2.45
    box(s, x, y, 2.3, 2.3, PINK if grp == 'kid' else SKY, None)
    img(s, f'games/demo/{ic}.png', x+0.6, y+0.15, w=1.1)
    text(s, x+0.05, y+1.3, 2.2, 0.45, nm, 13.5, INK, True, PP_ALIGN.CENTER)
    text(s, x+0.05, y+1.75, 2.2, 0.4, f'{g} · {age}', 11.5, SUB, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.85, 12, 0.3, '분홍 = 아이와 함께(광고 최소·결제 없음) · 하늘 = 어른을 위한', 11, SUB)

# 9 게임 화면
s = p.slides.add_slide(BL); bg(s); title(s, 'PLAY', '이렇게 놉니다'); foot(s)
shots = [('notes/snapshots/2026-10-07_merge_home.png', '도형 합치기'), ('notes/snapshots/2026-10-08_spot_home.png', '다른 그림 찾기'), ('notes/snapshots/2026-10-08_tile_home.png', '짝 맞추기'),
         ('notes/snapshots/saga_block_start.png', '이야기 지도'), ('notes/snapshots/2026-10-07_color_home.png', '막둥이 색칠북'), ('notes/snapshots/2026-10-08_day_home.png', '막둥이의 하루')]
for k, (f, nm) in enumerate(shots):
    x = 0.6 + k*2.05
    box(s, x, 1.55, 1.9, 4.95, WHITE, LINE, 0.12)
    img(s, f, x+0.08, 1.63, fit=(1.74, 4.3))
    text(s, x, 6.0, 1.9, 0.4, nm, 12.5, INK, True, PP_ALIGN.CENTER)

# 10 공통 시스템
s = p.slides.add_slide(BL); bg(s); title(s, 'ONE HOUSE', '열 개 게임을 하나로 묶는 장치', '어느 게임에서 얻든 같은 「우리 집」 · 같은 옛이야기 · 같은 배역'); foot(s)
album = crop('notes/snapshots/ip_sheet.png', (14, 14, 314, 614), 'album.png'); stamp = crop('notes/snapshots/ip_sheet.png', (642, 14, 942, 614), 'stamp.png')
cameo = crop('notes/snapshots/ip_sheet.png', (328, 660, 628, 1260), 'cameo.png'); rank = crop('notes/snapshots/rank_sheet.png', (0, 0, 300, 600), 'rank.png')
sysm = [('notes/snapshots/2026-10-07_merge_house.png', '우리 집', '모은 가구로 꾸미는 하나의 집'), (album, '이야기 카드', '옛이야기 17편, 여러 게임에서 모아 완성'),
        (stamp, '도장책', '게임마다 도장, 모으면 문패 장식'), (cameo, '카메오 도감', '다른 게임 가족이 깜짝 방문'), (rank, '세 가족 순위표', '서버 없이 가족과 겨루기')]
for k, (f, nm, d) in enumerate(sysm):
    x = 0.6 + k*2.45
    box(s, x, 1.95, 2.3, 4.85)
    img(s, f, x+0.1, 2.05, fit=(2.1, 3.4))
    text(s, x+0.1, 5.55, 2.1, 0.4, nm, 15, [NEMO, SEMO, DONG, NEMO, SEMO][k], True, PP_ALIGN.CENTER)
    text(s, x+0.1, 5.95, 2.1, 0.8, d, 11.5, INK, align=PP_ALIGN.CENTER)

# 11 아이와 부모가 안심
s = p.slides.add_slide(BL); bg(s); title(s, 'SAFE & KIND', '압박 없는 게임', '하트·대기시간·뽑기 없음 — 40~60대와 아이 부모가 안심하는 브랜드', DONG); foot(s)
img(s, 'notes/snapshots/2026-10-07_color_home.png', 0.6, 1.9, h=4.9)
img(s, 'notes/snapshots/2026-10-08_day_home.png', 2.95, 1.9, h=4.9)
pts = ['실패해도 벌 없이 바로 「한 번 더」, 되돌리기·힌트', '옛이야기 + 도형 가족 배역 — 설명 없이 통하는 이야기',
       '유아 앱: 글자 없는 놀이, 보호자 잠금, 쉬어요 알림, 결제·알림·추적 없음', '광고는 덜 거슬리게: 보상형은 선택, 판 사이 광고는 빈도 제한',
       '큰 글씨·움직임 줄이기·색 구분 보조 · 오프라인으로 동작']
text(s, 5.6, 2.1, 7.1, 4.6, ['· ' + t for t in pts], 15.5, INK, lsp=1.5)

# 12 다음 (원정대·굿즈)
s = p.slides.add_slide(BL); bg(s); title(s, 'NEXT', '다음 이야기', None, NEMO); foot(s)
box(s, 0.6, 1.6, 7.2, 5.2)
text(s, 0.9, 1.8, 6.6, 0.5, '⑪ 막둥이 원정대 (가칭) — 아이용 방치형 RPG', 19, NEMO, True)
for k, f in enumerate(['assets/nemo_kids/baby.png', 'assets/wife/joy.png', 'assets/dong_dad/calm.png']): img(s, f, 1.1 + k*2.1, 2.5, fit=(1.9, 2.0))
text(s, 0.9, 4.65, 6.6, 2.1, ['싸우지 않고 「도와주는」 RPG — 옛이야기 17나라를 걸으며 곤란한 이웃을 돕고 마을을 다시 세운다.', '동료는 뽑기 없이 합류, 꺼 두어도 천천히 진행, 하루 3분.'], 13.5, INK, lsp=1.35)
box(s, 8.1, 1.6, 4.6, 5.2)
text(s, 8.4, 1.8, 4.0, 0.5, '굿즈 · 협업', 19, SEMO, True)
for k, f in enumerate(['assets/wife/love.png', 'assets/nemo_mom/joy.png']): img(s, f, 8.5 + k*2.0, 2.5, fit=(1.8, 1.8))
text(s, 8.4, 4.5, 4.0, 2.2, ['스티커 · 노트 · 키링', '상품에 「우리 집」 아이템 코드 동봉 — 실물이 게임 속 집으로', '이모티콘 · 캐릭터 협업'], 13.5, INK, lsp=1.35)

# 13 연락
s = p.slides.add_slide(BL); bg(s)
for i, f in enumerate(FAM3): img(s, f, 0.9 + i*1.75, 1.0, h=2.2)
text(s, 0.8, 3.5, 6.5, 0.9, '당신은 어느 도형인가요?', 32, INK, True)
chan = ['게임 체험: ' + PLAY, '네이버 블로그(웹툰): blog.naver.com/gihahakf', '유튜브: @geometric.family']
if EMAIL: chan.append('문의: ' + EMAIL)
text(s, 0.8, 4.5, 7.5, 2.0, chan, 14, SUB, lsp=1.45)
qr = os.path.join(TMP, 'qr.png'); segno.make(PLAY, error='m').save(qr, scale=10, border=2)
box(s, 9.0, 1.4, 3.3, 3.9, WHITE, None)
img(s, qr, 9.25, 1.6, w=2.8)
text(s, 9.0, 4.55, 3.3, 0.5, '휴대폰으로 바로 해 보기', 14, INK, True, PP_ALIGN.CENTER)

p.save(OUT); print(OUT)
