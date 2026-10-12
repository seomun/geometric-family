from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import os

R = r'C:/projects/geometric-family/'
OUT = r'C:/Users/gurud/Downloads/기하학가족_사업소개_발표.pptx'
W, H = 12192000, 6858000
IN = 914400
CREAM = RGBColor(0xFF, 0xF6, 0xEA); INK = RGBColor(0x4A, 0x3A, 0x30); SUB = RGBColor(0x8A, 0x75, 0x66)
NEMO = RGBColor(0xE0, 0x7B, 0x39); SEMO = RGBColor(0x3A, 0x8F, 0xD6); DONG = RGBColor(0x3B, 0x4A, 0x7A)
YEL = RGBColor(0xF6, 0xDE, 0x5C); WHITE = RGBColor(0xFF, 0xFF, 0xFF); LINE = RGBColor(0xE8, 0xD9, 0xC6)
FONT = '맑은 고딕'

p = Presentation(); p.slide_width = W; p.slide_height = H
BL = p.slide_layouts[6]

def bg(s, c=CREAM):
    f = s.background.fill; f.solid(); f.fore_color.rgb = c

def text(s, x, y, w, h, t, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, lsp=1.15):
    tb = s.shapes.add_textbox(int(x*IN), int(y*IN), int(w*IN), int(h*IN)); tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    lines = t if isinstance(t, list) else [t]
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.line_spacing = lsp
        segs = ln.split('**')
        for j, seg in enumerate(segs):
            if not seg: continue
            r = para.add_run(); r.text = seg; r.font.size = Pt(size); r.font.name = FONT
            r.font.bold = bold or (j % 2 == 1); r.font.color.rgb = color
    return tb

def box(s, x, y, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    b = s.shapes.add_shape(shape, int(x*IN), int(y*IN), int(w*IN), int(h*IN))
    b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None: b.line.fill.background()
    else: b.line.color.rgb = line; b.line.width = Pt(1)
    b.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0] = 0.12
    return b

def img(s, path, x, y, w=None, h=None):
    path = R + path
    iw, ih = Image.open(path).size
    if w and not h: h = w * ih / iw
    if h and not w: w = h * iw / ih
    return s.shapes.add_picture(path, int(x*IN), int(y*IN), int(w*IN), int(h*IN))

def title(s, n, t, sub=None):
    text(s, 0.6, 0.35, 0.6, 0.4, f'{n:02d}', 14, SUB, True)
    text(s, 0.6, 0.62, 12, 0.7, t, 30, INK, True)
    if sub: text(s, 0.6, 1.3, 12, 0.4, sub, 15, SUB)
    ln = s.shapes.add_connector(1, int(0.6*IN), int(1.78*IN), int(12.73*IN), int(1.78*IN)); ln.line.color.rgb = LINE; ln.line.width = Pt(1.5)

def foot(s):
    text(s, 0.6, 7.05, 8, 0.3, '기하학 가족', 10, SUB)

FAM = [('assets/nemo_dad/good.png', NEMO, '네모', '많고 벅차지만 북적이는 행복'),
       ('assets/wife/good.png', SEMO, '세모', '극단 — 티격태격과 달달함'),
       ('assets/dong_dad/good.png', DONG, '동그라미', '완벽한데 어딘가 비어 있음')]

# 1 표지
s = p.slides.add_slide(BL); bg(s)
for i, (pth, c, nm, _) in enumerate(FAM):
    img(s, pth, 7.0 + i*2.0, 2.2, h=2.6)
text(s, 0.8, 1.6, 6.2, 1.0, '기하학 가족', 54, INK, True)
text(s, 0.8, 2.75, 6.2, 1.2, ['네모·세모·동그라미 세 도형 가족이', '같은 일을 각자 다르게 겪는 이야기 IP'], 20, SUB)
box(s, 0.8, 4.35, 5.6, 0.62, YEL, None)
text(s, 0.8, 4.35, 5.6, 0.62, '웹툰 · 쇼츠로 알리고, 게임 · 굿즈로 법니다', 18, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 0.8, 5.7, 8, 0.9, ['대표 네모 (1인 창작자 · 예비 재창업자)'], 15, SUB)

# 2 대표·창업 동기
s = p.slides.add_slide(BL); bg(s); title(s, 2, '왜 시작했나', '게임사를 접고, 캐릭터와 이야기로 다시 시작합니다'); foot(s)
items = [('2026년 7월', '운영하던 게임 회사를 접었다. AI 로 혼자도 게임을 만드는 시대 — 작은 게임은 쏟아지고, 잘 만드는 것만으로는 눈에 띄기 어렵다.'),
         ('깨달음', '노래 한 곡보다 가수와 스토리가 소비되는 시대. 오래가는 것은 캐릭터와 세계관이다.'),
         ('출발점', '오래전 내가 써 둔 웹툰 — 가족이 모두 도형인 이야기. 장례식장에 다녀와 쓴 첫 장.'),
         ('공식', '콘텐츠는 소비되지만, 경험은 쌓인다. 예전엔 여러 사람이 필요했지만, 지금은 AI 와 1인이 할 수 있다.')]
for i, (k, v) in enumerate(items):
    y = 2.1 + i*1.15
    box(s, 0.6, y, 2.2, 0.9, [NEMO, SEMO, DONG, YEL][i], None)
    text(s, 0.6, y, 2.2, 0.9, k, 17, WHITE if i < 3 else INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, 3.05, y+0.05, 9.6, 0.85, v, 16, INK, anchor=MSO_ANCHOR.MIDDLE)

# 3 세 가족 = 하나의 사건, 세 가지 반응
s = p.slides.add_slide(BL); bg(s); title(s, 3, '하나의 사건, 세 가족의 반응', '50대 동창 세 남자와 그 가족들 — 같은 일을 각자의 방식으로 겪는다'); foot(s)
for i, (pth, c, nm, d) in enumerate(FAM):
    x = 0.6 + i*4.1
    box(s, x, 2.05, 3.85, 4.35)
    img(s, pth, x + (3.85-2.0)/2, 2.25, h=2.3)
    text(s, x, 4.7, 3.85, 0.5, nm, 24, c, True, PP_ALIGN.CENTER)
    text(s, x+0.2, 5.25, 3.45, 1.2, d, 15, INK, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.55, 12, 0.35, '매 화 끝 질문 — 「당신은 어느 도형인가요?」 댓글과 공유가 IP 를 퍼뜨립니다', 14, SUB, True, PP_ALIGN.CENTER)

# 4 제품 라인업
s = p.slides.add_slide(BL); bg(s); title(s, 4, '하나의 세계관, 여러 매체', '웹툰·쇼츠는 입구, 게임·굿즈는 수익'); foot(s)
cards = [('웹툰', '공감툰 · 네이버 블로그 연재', '2화 제작, 원고 15편'),
         ('쇼츠', '통계 + 세 가족 반응, 작가 나레이션', '3편 · 유튜브·릴스·틱톡'),
         ('게임 10종', '유아 3 + 성인 7, 하나의 「우리 집」 공유', '출시 후보 완성 · 웹 시험판'),
         ('정보 블로그', '「짠내나는 경제학」 퇴직 후 생활 통계', '티스토리 준비'),
         ('굿즈', '스티커·노트 + 게임 아이템 코드', '반응 확인 후')]
for i, (a, b, c) in enumerate(cards):
    x = 0.6 + i*2.45
    box(s, x, 2.1, 2.3, 3.0)
    text(s, x, 2.3, 2.3, 0.5, a, 20, [NEMO, SEMO, DONG, NEMO, SEMO][i], True, PP_ALIGN.CENTER)
    text(s, x+0.15, 2.95, 2.0, 1.3, b, 13.5, INK, align=PP_ALIGN.CENTER)
    box(s, x+0.15, 4.35, 2.0, 0.55, CREAM, None)
    text(s, x+0.15, 4.35, 2.0, 0.55, c, 11.5, SUB, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

flow = ['웹툰 · 쇼츠로 만남', '「어느 도형인가요?」 댓글·공유', '게임에서 놀고 「우리 집」 꾸미기', '굿즈 · 아이템 코드']
text(s, 0.6, 5.3, 12, 0.4, '팬이 되는 길', 15, SUB, True)
for i, fl in enumerate(flow):
    x = 0.6 + i*3.08
    box(s, x, 5.75, 2.75, 0.9, [NEMO, SEMO, DONG, YEL][i], None)
    text(s, x+0.1, 5.75, 2.55, 0.9, fl, 14, WHITE if i < 3 else INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if i < 3: text(s, x+2.75, 5.75, 0.33, 0.9, '›', 24, SUB, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# 5 게임 10종
s = p.slides.add_slide(BL); bg(s); title(s, 5, '게임 10종 — 출시 후보 완성', '장르는 다르고, 화면·소리·「우리 집」은 하나'); foot(s)
shots = ['notes/snapshots/2026-10-07_block_home.png', 'notes/snapshots/2026-10-08_tile_home.png', 'notes/snapshots/saga_block_start.png', 'notes/snapshots/2026-10-08_day_home.png']
for i, sh in enumerate(shots):
    img(s, sh, 0.6 + i*1.85, 2.05, h=3.3)
lst = ['유아: 놀이터 · 막둥이 색칠북 · 막둥이의 하루', '성인: 세 가족 식탁(방치) · 도형 합치기 · 어느 도형?(테스트) · 다른 그림 찾기 · 도형 블록 · 정리의 달인 · 세 가족 짝 맞추기',
       '게임마다 판 60~140개, 옛이야기 12편 + 도형 가족 배역', '이야기 지도 · 이벤트 달력 · 순위표 · 출석 · 집 꾸미기', '광고·결제 모듈, 아동 정책 점검 완료']
text(s, 8.2, 2.05, 4.5, 4.6, ['· ' + l for l in lst], 14, INK, lsp=1.3)
text(s, 0.6, 5.55, 7.5, 0.4, '웹 시험판: seomun.github.io/geometric-family/play', 13, SEMO, True)

# 6 차별점
s = p.slides.add_slide(BL); bg(s); title(s, 6, '무엇이 다른가'); foot(s)
diff = [('「우리 집」 공유', '어느 게임에서 얻은 가구든 같은 집에 놓인다 → 게임을 하나 더 깔 이유'),
        ('아는 이야기 + 우리 배역', '아기돼지 삼형제·흥부전 … 설명 없이 통하는 옛이야기에 도형 가족이 반복 등장'),
        ('압박 없는 설계', '하트·대기시간·뽑기 없음 → 40~60대와 아이 부모가 안심하는 브랜드'),
        ('1인 + AI 제작 체계', '세계관 문서 → AI 가 설계·코딩·검수 → 10종을 약 4주에, 고정비 거의 0')]
for i, (a, b) in enumerate(diff):
    x = 0.6 + (i % 2)*6.1; y = 2.1 + (i//2)*2.35
    box(s, x, y, 5.95, 2.15)
    text(s, x+0.3, y+0.25, 5.4, 0.5, a, 20, [NEMO, SEMO, DONG, NEMO][i], True)
    text(s, x+0.3, y+0.85, 5.4, 1.2, b, 15, INK)

# 7 성과
s = p.slides.add_slide(BL); bg(s); title(s, 7, '지금까지 한 것 (2026년 9~10월, 약 4주)'); foot(s)
nums = [('10', '게임 (출시 후보)'), ('10', '캐릭터 정본'), ('2', '웹툰 화'), ('3', '쇼츠')]
for i, (n, l) in enumerate(nums):
    x = 0.6 + i*3.05
    box(s, x, 2.05, 2.85, 1.9, WHITE)
    text(s, x, 2.15, 2.85, 1.0, n, 48, [NEMO, SEMO, DONG, NEMO][i], True, PP_ALIGN.CENTER)
    text(s, x, 3.2, 2.85, 0.5, l, 15, SUB, True, PP_ALIGN.CENTER)
done = ['세계관 백서 · 캐릭터 10명 정본 시트 · 상표 출원 준비(이름·캐릭터 그림)', '네이버 블로그 · 유튜브 · 인스타 채널 개설, 웹툰·쇼츠 연재 시작',
        '게임 10종: 판 1,000여 개, 공통 화면·소리·집·이야기 지도·광고·결제, 자동 점검 통과', '출시 문서(스토어 문안·개인정보처리방침·아동 정책 점검표), 최저가 서버 계획']
text(s, 0.6, 4.3, 7.4, 2.6, ['· ' + d for d in done], 14, INK, lsp=1.35)
for k, c in enumerate(['c01', 'c08', 'c09']):
    img(s, f'content/blog/ep01_short_cuts/{c}.png', 8.3 + k*1.5, 4.25, w=1.4)

# 8 수익 모델
s = p.slides.add_slide(BL); bg(s); title(s, 8, '어떻게 버는가'); foot(s)
rev = [('게임 광고', '성인: 보상형 + 판 사이(빈도 제한) + 배너 / 유아: 아동용 배너만', '출시 즉시'),
       ('게임 결제', '광고 제거 1회 구매 · 꾸미기 소품 팩(확정 지급)', '출시 후'),
       ('퍼블리싱', '게임 10종 퍼블리셔 계약 — 출시·마케팅·서버비 분담', '협의 중'),
       ('콘텐츠', '블로그 광고 · 쇼츠 수익', '조건 충족 후'),
       ('굿즈·협업', '스티커·노트 + 게임 아이템 코드', '반응 확인 후')]
for i, (a, b, c) in enumerate(rev):
    y = 2.05 + i*0.9
    box(s, 0.6, y, 12.1, 0.78)
    text(s, 0.9, y, 2.4, 0.78, a, 17, [NEMO, SEMO, DONG, NEMO, SEMO][i], True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.3, y, 7.0, 0.78, b, 14.5, INK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 10.4, y, 2.1, 0.78, c, 13, SUB, True, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.6, 12, 0.4, '가정: 성인 앱 하루 이용자 1,000명 → 광고 월 약 $150~600 (업계 평균 범위의 가정, 출시 후 실제 지표로 갱신)', 12.5, SUB)

# 9 로드맵
s = p.slides.add_slide(BL); bg(s); title(s, 9, '앞으로 12개월'); foot(s)
road = [('1~3개월', ['그림·음악 교체(AI 툴 한 달)', '상표 출원 · 사업자 등록', '게임 10종 퍼블리셔 납품·출시', '웹툰 주 1화 · 쇼츠 주 2~3편']),
        ('4~6개월', ['출시 지표 보고 개선', '광고 제거 · 소품 팩 판매', '굿즈 소량 시험(POD)']),
        ('7~12개월', ['인기 게임 판 추가 · 시즌 이벤트', '웹툰 해외판(영·일)', '다음 게임 묶음', '재창업 지원 신청(매출 근거)'])]
for i, (a, items) in enumerate(road):
    x = 0.6 + i*4.1
    box(s, x, 2.05, 3.85, 0.7, [NEMO, SEMO, DONG][i], None)
    text(s, x, 2.05, 3.85, 0.7, a, 20, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    box(s, x, 2.85, 3.85, 2.5)
    text(s, x+0.3, 3.05, 3.3, 2.2, ['· ' + t for t in items], 15, INK, lsp=1.45)

# 10 입주 활용 계획
s = p.slides.add_slide(BL); bg(s); title(s, 10, '입주하면 이렇게 쓰겠습니다', '집 밖의 작업실 = 1인 기업이 회사가 되는 첫 자리'); foot(s)
use = [('제작 공간', '웹툰·게임 제작, 쇼츠 나레이션 녹음(조용한 공간)'),
       ('파트너 미팅', '퍼블리셔·굿즈 제작처 미팅을 회의실에서'),
       ('멘토링 활용', '세무·법률(상표·저작권·AI 저작물), 마케팅 자문'),
       ('사업화 연계', '사업자 등록, 지원사업·투자 연계, 매출 근거 마련')]
for i, (a, b) in enumerate(use):
    x = 0.6 + (i % 2)*6.1; y = 2.1 + (i//2)*2.1
    box(s, x, y, 5.95, 1.9)
    text(s, x+0.3, y+0.25, 5.4, 0.5, a, 20, [NEMO, SEMO, DONG, NEMO][i], True)
    text(s, x+0.3, y+0.85, 5.4, 0.9, b, 15, INK)
text(s, 0.6, 6.45, 12, 0.4, '목표: 1년 안에 게임 출시 · 첫 매출 · 상표 등록', 14, INK, True)

# 11 마무리
s = p.slides.add_slide(BL); bg(s)
for i, (pth, c, nm, _) in enumerate(FAM):
    img(s, pth, 3.4 + i*2.3, 1.2, h=2.5)
text(s, 0.6, 4.1, 12.1, 0.9, '당신은 어느 도형인가요?', 40, INK, True, PP_ALIGN.CENTER)
text(s, 0.6, 5.1, 12.1, 0.9, ['콘텐츠는 소비되지만, 경험은 쌓입니다.', '기하학 가족이 회사가 되겠습니다. 감사합니다.'], 17, SUB, align=PP_ALIGN.CENTER)

p.save(OUT); print(OUT)
